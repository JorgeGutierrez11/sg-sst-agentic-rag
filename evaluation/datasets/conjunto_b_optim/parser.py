"""
Convierte la hoja "Dataset Conjunto B" de Diseño.xlsx a JSON.
Excluye la columna "Estado". Detecta encabezados automáticamente
(no depende de filas/columnas fijas).

Uso: python parser.py
"""
from pathlib import Path
import json
import openpyxl

XLSX_PATH = Path(__file__).resolve().parents[1] / "Diseño.xlsx"
SHEET_NAME = "Dataset Conjunto B"
EXCLUDE_COLS = {"Estado"}
OUTPUT_PATH = Path(__file__).resolve().parent / "conjunto_b_optim.json"


def find_header_row(ws):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        values = [c.value for c in row]
        if any(v == "ID" for v in values):
            return row[0].row
    raise ValueError("No se encontró fila de encabezados con 'ID'.")


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Hoja '{SHEET_NAME}' no existe. Hojas disponibles: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    header_row_num = find_header_row(ws)
    header_cells = ws[header_row_num]

    col_map = {}
    for cell in header_cells:
        if cell.value is None:
            continue
        if cell.value in EXCLUDE_COLS:
            continue
        col_map[cell.column] = cell.value

    records = []
    for row in ws.iter_rows(min_row=header_row_num + 1):
        row_by_col = {c.column: c.value for c in row}
        id_col = next((col for col, name in col_map.items() if name == "ID"), None)
        if id_col is not None and not row_by_col.get(id_col):
            continue
        record = {name: row_by_col.get(col) for col, name in col_map.items()}
        records.append(record)

    OUTPUT_PATH.write_text(
        json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"{len(records)} registros -> {OUTPUT_PATH}")


if __name__ == "__main__":
    main()